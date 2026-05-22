import openpyxl
import copy
from pptx import Presentation

def load_names_from_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    names = []
    # Skip header
    for row in sheet.iter_rows(min_row=2, max_col=1):
        cell = row[0]
        if cell.value:
            val = str(cell.value).strip()
            # Skip bold headers (e.g. BACHELOR OF..., Emphasis in...)
            if cell.font and cell.font.bold:
                print(f"Skipping program/section heading: '{val}'")
                continue
            names.append(val)
    return names

def create_lowerthird_ppt(names, template_file='template.pptx', output_file='lowerthirds.pptx'):
    if not names:
        print("No student names loaded. Exiting.")
        return

    # Load presentation
    prs = Presentation(template_file)
    template_slide = prs.slides[0]
    
    # Store shapes XML and picture relationship mappings from the template slide
    template_shapes_xml = []
    template_rels = {}
    for shape in template_slide.shapes:
        shape_el = copy.deepcopy(shape.element)
        template_shapes_xml.append((shape.shape_type, shape_el))
        
        if shape.shape_type == 13: # PICTURE
            try:
                blip_tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
                blip = shape.element.find(f'.//{blip_tag}')
                if blip is not None:
                    rId_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                    old_rId = blip.get(rId_attr)
                    template_rels[old_rId] = template_slide.part.rels[old_rId]
            except Exception as e:
                print(f"Failed to cache picture relationship: {e}")

    print(f"Generating {len(names)} slides...")

    # Process remaining names by creating new slides and copying template shapes
    for i, name in enumerate(names[1:], start=1):
        new_slide = prs.slides.add_slide(template_slide.slide_layout)
        
        # Copy shapes XML
        for shape_type, shape_el in template_shapes_xml:
            new_el = copy.deepcopy(shape_el)
            
            # Replace placeholder in XML directly to ensure persistence on save
            t_tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}t'
            for t_node in new_el.findall(f'.//{t_tag}'):
                if t_node.text and "Name" in t_node.text:
                    t_node.text = t_node.text.replace("Name", name)
            
            new_slide.shapes._spTree.append(new_el)
            
            # Map picture relationships
            if shape_type == 13: # PICTURE
                try:
                    blip_tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
                    blip = new_el.find(f'.//{blip_tag}')
                    if blip is not None:
                        rId_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
                        old_rId = blip.get(rId_attr)
                        rel = template_rels.get(old_rId)
                        if rel is not None:
                            new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
                            blip.set(rId_attr, new_rId)
                except Exception as e:
                    print(f"Failed to map picture relationship on slide {i}: {e}")

    # Process first name in-place on the original slide 0 at the very end
    slide0 = prs.slides[0]
    replaced_first = False
    for shape in slide0.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if "Name" in run.text:
                        run.text = run.text.replace("Name", names[0])
                        replaced_first = True
    if replaced_first:
        print(f"Slide 0 name replaced: '{names[0]}'")
                            
    # Save the output presentation
    prs.save(output_file)
    print(f"Success: {output_file} created with {len(names)} slides.")

if __name__ == "__main__":
    names = load_names_from_excel('names.xlsx')
    create_lowerthird_ppt(names)
